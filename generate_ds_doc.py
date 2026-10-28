"""
Script to generate the comprehensive CardioSense Complete Data Science Specification document in .docx format.
Updated to reflect the dual-cohort architecture:
- Primary Focus: Indian Hospital Cardiovascular Disease Dataset (Mendeley Data DOI: 10.17632/dzz48mvjht.1, Doppala & Bhattacharyya, N=1,000)
- Benchmark Baseline: UCI Cleveland Heart Disease Cohort (Janosi et al., 1988, N=303)
- Research Context: ICMR-INDIAB study and Indian Coronary Artery Disease Dataset (ICADD)
- Zero-Cholesterol Sentinel Imputation (53 patients with serumcholestrol == 0 coerced to NaN)
- Stratified 5-Fold Cross-Validation, Hyperparameter Tuning & Leakage-Safe ColumnTransformer Pipelines
- Real Measured Benchmark Metrics (XGBoost 99.50% Acc, 100.0% Recall, 0 False Negatives on India cohort)
- Game-Theoretic SHAP Attributions (Summary Beeswarm, Feature Importance Bar, and Waterfall Plots)
- Production Serving Architecture: FastAPI (/predict/india, /predict/uci), Streamlit Pastel UI, Next.js Portal
"""

import sys
import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

PROJECT_ROOT = Path(r"D:\Aditya's Folder\DS CP\CardioSense")
FIGURES_DIR_INDIA = PROJECT_ROOT / "reports" / "figures" / "india"
FIGURES_DIR_UCI = PROJECT_ROOT / "reports" / "figures" / "uci"
FIGURES_DIR_ROOT = PROJECT_ROOT / "reports" / "figures"
OUTPUT_DOCX = PROJECT_ROOT / "CardioSense_Complete_Data_Science_Specification.docx"

# Color Palette: Deep Slate Navy, Medical Steel Blue, Vibrant Cyan, Charcoal Body, Off-White
PRIMARY_COLOR = RGBColor(26, 54, 93)       # #1A365D - Deep Navy
SECONDARY_COLOR = RGBColor(43, 108, 176)   # #2B6CB0 - Medical Steel Blue
ACCENT_COLOR = RGBColor(49, 130, 206)      # #3182CE - Vibrant Cyan/Blue
TEXT_DARK = RGBColor(45, 55, 72)           # #2D3748 - Charcoal
MUTED_TEXT = RGBColor(113, 128, 150)       # #718096 - Slate Muted
CALLOUT_BG = "F7FAFC"                      # Light Ice Grey
HEADER_BG = "1A365D"                       # Dark Navy Header
ZEBRA_BG = "EDF2F7"                        # Light row tint
BORDER_COLOR = "CBD5E0"                    # Clean border grey

def set_cell_shading(cell, color_hex):
    """Apply background color to table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell internal padding in twips (1/20 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    """Set subtle border on entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_styled(doc, text, level):
    """Add a heading with exact corporate typography and spacing."""
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    run = p.runs[0]
    if level == 1:
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = ACCENT_COLOR
    return p

def add_paragraph_styled(doc, text="", bold_prefix="", italic=False, space_after=5):
    """Add styled body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = TEXT_DARK
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10.5)
        r_text.font.italic = italic
        r_text.font.color.rgb = TEXT_DARK
    return p

def add_callout(doc, title, text, icon="ℹ️"):
    """Create a styled callout box table."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E0", sz="6", val="single")
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"{icon} {title}\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR
    
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = TEXT_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_formula_box(doc, title, formula_text, explanation_text):
    """Create a callout box specifically formatted for mathematical equations."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="2B6CB0", sz="8", val="single")
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"Mathematical Formulation: {title}\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = SECONDARY_COLOR
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(4)
    r_eq = p2.add_run(formula_text)
    r_eq.font.name = "Consolas"
    r_eq.font.size = Pt(10)
    r_eq.font.bold = True
    r_eq.font.color.rgb = PRIMARY_COLOR
    
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(2)
    r_exp = p3.add_run(explanation_text)
    r_exp.font.name = "Calibri"
    r_exp.font.size = Pt(9.5)
    r_exp.font.italic = True
    r_exp.font.color.rgb = TEXT_DARK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_with_caption(doc, image_name, caption, cohort="india", width_in_inches=5.8):
    """Add a figure centered with caption, checking cohort directories first."""
    if cohort == "india":
        img_path = FIGURES_DIR_INDIA / image_name
    elif cohort == "uci":
        img_path = FIGURES_DIR_UCI / image_name
    else:
        img_path = FIGURES_DIR_ROOT / image_name

    if not img_path.exists():
        img_path = FIGURES_DIR_ROOT / image_name
    if not img_path.exists():
        p = doc.add_paragraph(f"[Image Missing: {image_name}]")
        p.runs[0].font.italic = True
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(str(img_path), width=Inches(width_in_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(8)
    run_cap = p_cap.add_run(f"Figure: {caption}")
    run_cap.font.name = "Calibri"
    run_cap.font.size = Pt(9)
    run_cap.font.italic = True
    run_cap.font.color.rgb = MUTED_TEXT

def create_styled_table(doc, headers, data, col_widths=None):
    """Render a high-end styled table with dark header and zebra stripes."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=BORDER_COLOR, sz="4", val="single")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        set_cell_shading(hdr_cells[i], HEADER_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = ZEBRA_BG if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(9)
            p.runs[0].font.color.rgb = TEXT_DARK
            
    # Widths
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def build_document():
    doc = docx.Document()
    
    # Page setup - 1 inch margins all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header & Footer setup for subsequent pages
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = "CardioSense: India-Focused Cardiovascular Disease ML Specification & Research Monograph"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.runs[0].font.name = "Calibri"
        p_hdr.runs[0].font.size = Pt(8.5)
        p_hdr.runs[0].font.color.rgb = MUTED_TEXT
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.text = "CardioSense Academic & Scientific Report | Confidential & Educational Use Only"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_ftr.runs[0].font.name = "Calibri"
        p_ftr.runs[0].font.size = Pt(8.5)
        p_ftr.runs[0].font.color.rgb = MUTED_TEXT

    # ==========================================
    # COVER / TITLE BLOCK
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("CardioSense\n")
    r_t1.font.name = "Calibri"
    r_t1.font.size = Pt(28)
    r_t1.font.bold = True
    r_t1.font.color.rgb = PRIMARY_COLOR
    
    r_sub = p_title.add_run("Dual-Cohort Machine Learning Lifecycle, Clinical Rationale, Methodology & Research Monograph")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = SECONDARY_COLOR
    
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_before = Pt(4)
    p_tag.paragraph_format.space_after = Pt(24)
    r_tag = p_tag.add_run("Primary Focus: Indian Hospital Cohort (N=1,000) & Canonical UCI Cleveland Benchmark (N=303)")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(11)
    r_tag.font.italic = True
    r_tag.font.color.rgb = MUTED_TEXT

    # Metadata Table
    meta_headers = ["Attribute", "Specification / Value"]
    meta_data = [
        ["Project Repository", "https://github.com/Adi3595/CardioSense"],
        ["Primary Author & Lead", "Aditya (Adi3595)"],
        ["Primary Research Cohort", "Indian Hospital Cohort (Mendeley Data DOI: 10.17632/dzz48mvjht.1)"],
        ["Indian Cohort Dimensions", "1,000 Patient Records × 14 Raw Attributes (12 Predictors + 1 ID + 1 Target)"],
        ["Primary Checksum", "SHA256: 9376367dde41e79e432d6402ea5d9e44223fc318ff54ed09966d721923947d1e"],
        ["Benchmark Cohort", "UCI Cleveland Clinic Foundation Database (N=303, Janosi et al., 1988)"],
        ["Benchmark Checksum", "SHA256: a74b7efa387bc9d108d7d0115d831fe9b414b29ae7124f331b622b4efa0427c8"],
        ["Champion Model (India)", "XGBoost Classifier (Test Acc: 99.50%, Recall: 100.0%, 0 False Negatives, ROC-AUC: 0.9998)"],
        ["Champion Model (UCI)", "Random Forest Classifier (CV ROC-AUC: 0.9038 ± 0.0238, Test ROC-AUC: 0.9502)"],
        ["Explainability Engine", "TreeExplainer via Shapley Additive exPlanations (SHAP)"],
        ["Deployment Stack", "Python 3.12, Scikit-Learn 1.9, XGBoost, FastAPI 0.116, Streamlit 1.42, Next.js 16"]
    ]
    create_styled_table(doc, meta_headers, meta_data, [2.2, 4.3])
    
    add_callout(
        doc,
        "MANDATORY MEDICAL & NON-DIAGNOSTIC DISCLAIMER",
        "CardioSense is an academic research, machine learning engineering, and educational prototype. "
        "It is designed solely to study clinical feature attribution, leakage-free MLOps pipelines, and diagnostic model trade-offs. "
        "It is NOT clinically validated, NOT certified as medical device software (FDA / CE-MDR), and MUST NOT be used "
        "for actual clinical triage, patient diagnosis, prescription, or treatment planning. All clinical decisions require "
        "a licensed physician and standard diagnostic angiography.",
        icon="⚠️"
    )

    doc.add_page_break()

    # ==========================================
    # SECTION 1: THE RESEARCH WHY & CLINICAL FORMULATION
    # ==========================================
    add_header_styled(doc, "1. Executive Summary & Clinical Problem Formulation", level=1)
    
    add_paragraph_styled(
        doc,
        "Cardiovascular diseases (CVDs) represent the single leading cause of global morbidity and premature mortality, "
        "accounting for an estimated 17.9 million deaths annually according to the World Health Organization. "
        "In South Asia and India specifically, cardiovascular disease presents an unprecedented public health crisis. "
        "Epidemiological investigations, including the landmark Indian Council of Medical Research–India Diabetes (ICMR-INDIAB) "
        "study and the Indian Coronary Artery Disease Dataset (ICADD), demonstrate that individuals of South Asian ancestry "
        "develop coronary artery disease (CAD) 5 to 10 years earlier than Western populations, with a substantially higher "
        "incidence of multi-vessel diffuse disease and sudden cardiac arrest during their peak productive working years."
    )
    
    add_header_styled(doc, "1.1 The Clinical Diagnostic Dilemma: Invasive Angiography vs. Non-Invasive Screening", level=2)
    add_paragraph_styled(
        doc,
        "The definitive reference procedure for confirming obstructive CAD is invasive coronary catheter angiography. "
        "During catheterization, a radiopaque iodinated contrast agent is injected directly into coronary arteries via femoral or radial access. "
        "While definitive, invasive catheterization entails non-trivial procedural risks:"
    )
    
    cath_risks = [
        ("Vascular Complications: ", "Arterial dissection, pseudoaneurysms, retroperitoneal hemorrhage, and access-site hematomas."),
        ("Systemic & Renal Toxicity: ", "Contrast-induced acute kidney injury (nephropathy), especially prevalent in Indian diabetic populations."),
        ("Neurological Events: ", "Catheter-induced thromboembolic strokes or transient ischemic attacks (0.1–0.3% incidence)."),
        ("Resource & Financial Burden: ", "High procedural cost ($2,000–$15,000) and reliance on specialized catheterization laboratories, severely restricted in rural/semi-urban India."),
        ("High Negative Diagnostic Rate: ", "Historical registries indicate that up to 35% to 50% of elective catheterizations performed on moderately symptomatic patients reveal non-obstructive coronary arteries, subjecting healthy individuals to procedural risks.")
    ]
    for bold_pre, text in cath_risks:
        add_paragraph_styled(doc, text=text, bold_prefix="• " + bold_pre, space_after=3)

    add_header_styled(doc, "1.2 The Dual-Cohort Architecture & Scientific Research Questions", level=2)
    add_paragraph_styled(
        doc,
        "To reflect the specific risk dynamics of the Indian population while preserving comparability against Western benchmarks, "
        "CardioSense implements a dual-cohort design investigating five primary research questions:"
    )
    
    rqs = [
        ("RQ1 (Cohort-Specific Model Parity): ", "How do linear baselines (L2 Logistic Regression) compare against non-linear ensemble models (Random Forest, XGBoost) on the Indian hospital cohort (N=1,000) versus the Western UCI benchmark (N=303)?"),
        ("RQ2 (Sensitivity & False Negative Elimination): ", "Can gradient boosted architectures achieve zero or near-zero False Negatives on clinical test splits, effectively addressing the risk of fatal missed cardiac diagnoses?"),
        ("RQ3 (Feature Attribution Comparison via SHAP): ", "Which clinical predictors exert the strongest influence on Indian hospital predictions versus the Western Cleveland benchmark? Do non-invasive predictors (chest pain type, ST depression, exercise angina) dominate over traditional lipid profiles?"),
        ("RQ4 (Zero Data Leakage & Domain Hygiene): ", "Can domain-specific placeholders (such as 53 zero-cholesterol sentinels in the Indian dataset) be handled through leakage-safe training-fold median imputation without test-set contamination?"),
        ("RQ5 (Architectural Separation): ", "How do separate dataset-specific schemas and models prevent the confounding errors that occur when incompatible medical datasets are artificially merged?")
    ]
    for bold_pre, text in rqs:
        add_paragraph_styled(doc, text=text, bold_prefix=bold_pre, space_after=4)

    # ==========================================
    # SECTION 2: DATASET PROVENANCE, SCHEMAS & QUALITY AUDIT
    # ==========================================
    add_header_styled(doc, "2. Dual Dataset Provenance, Schemas & Quality Audit", level=1)
    
    add_paragraph_styled(
        doc,
        "CardioSense deliberately maintains two separate, unmerged datasets. Incompatible feature spaces are never concatenated "
        "to prevent spurious associations and cross-cohort distribution leakage."
    )
    
    add_header_styled(doc, "2.1 Primary Dataset: Indian Hospital Cohort", level=2)
    add_paragraph_styled(
        doc,
        "The primary dataset was acquired from a multi-specialty tertiary care hospital in India, published by Kasturi Doppala "
        "and Debnath Bhattacharyya (2020) on Mendeley Data (DOI: 10.17632/dzz48mvjht.1). It comprises 1,000 patient records "
        "with 14 original columns (1 patient identifier, 12 clinical predictors, and 1 binary disease target)."
    )

    audit_points_india = [
        ("Patient Identifier Dropping: ", "Column 'patientid' was dropped immediately upon ingestion to eliminate identification leakage."),
        ("Target Balance: ", "58.0% positive (580 CVD patients) and 42.0% negative (420 healthy controls), providing a rich supervised signal."),
        ("Younger Mean Age: ", "Mean cohort age is 49.2 ± 7.9 years (range 20–80 years), significantly younger than Western cohorts (54.4 years), confirming premature CAD demographics."),
        ("Elevated Resting Blood Pressure: ", "Mean systolic blood pressure is 151.7 ± 23.5 mm Hg, reflecting high chronic hypertension prevalence."),
        ("Zero-Cholesterol Clinical Sentinel: ", "53 patients (5.3%) had serum cholesterol recorded as exactly 0.0 mg/dL. In Indian acute hospital triage, patients admitted with acute coronary syndrome often undergo emergency catheterization before fasting lipid panels can be drawn. Treating 0 as a true biological value would severely distort modeling; these values were coerced to NaN and imputed using the training-fold median.")
    ]
    for bold_pre, text in audit_points_india:
        add_paragraph_styled(doc, text=text, bold_prefix="• " + bold_pre, space_after=3)

    add_header_styled(doc, "2.2 Indian Cohort Feature Dictionary", level=3)
    india_feat_headers = ["Feature Name", "Type", "Clinical Meaning", "Valid Range / Domain", "Preprocessing Pipeline"]
    india_feat_data = [
        ["age", "Numeric", "Age in completed years", "20 – 90 yrs", "Median Imputation + StandardScaler"],
        ["gender", "Categorical", "Biological sex (1=Male, 0=Female)", "{0, 1}", "Passthrough / OneHot"],
        ["chestpain", "Categorical", "Chest pain type (0=Typical, 1=Atypical, 2=Non-anginal, 3=Asymptomatic)", "{0, 1, 2, 3}", "OneHotEncoder(handle_unknown='ignore')"],
        ["restingBP", "Numeric", "Resting systolic blood pressure", "90 – 220 mm Hg", "Median Imputation + StandardScaler"],
        ["serumcholestrol", "Numeric", "Serum total cholesterol (zeros -> NaN)", "100 – 600 mg/dL", "Median Imputation + StandardScaler"],
        ["fastingbloodsugar", "Categorical", "Fasting blood sugar > 120 mg/dL", "{0, 1}", "Passthrough / OneHot"],
        ["restingrelectro", "Categorical", "Resting ECG (0=Normal, 1=ST abnormality, 2=LV hypertrophy)", "{0, 1, 2}", "OneHotEncoder(handle_unknown='ignore')"],
        ["maxheartrate", "Numeric", "Maximum heart rate achieved during exercise", "60 – 240 bpm", "Median Imputation + StandardScaler"],
        ["exerciseangia", "Categorical", "Exercise-induced angina (1=Yes, 0=No)", "{0, 1}", "Passthrough / OneHot"],
        ["oldpeak", "Numeric", "ST depression induced by exercise relative to rest", "0.0 – 6.5 mm", "Median Imputation + StandardScaler"],
        ["slope", "Categorical", "Slope of peak exercise ST segment (1=Up, 2=Flat, 3=Down)", "{1, 2, 3}", "OneHotEncoder(handle_unknown='ignore')"],
        ["noofmajorvessels", "Discrete", "Number of major vessels (0–3) colored by fluoroscopy", "{0, 1, 2, 3}", "OneHotEncoder(handle_unknown='ignore')"],
        ["target", "Binary", "Cardiovascular disease presence (1=Presence, 0=Absence)", "{0, 1}", "Supervised Binary Ground Truth"]
    ]
    create_styled_table(doc, india_feat_headers, india_feat_data, [1.4, 0.9, 2.0, 1.0, 1.2])

    add_header_styled(doc, "2.3 Benchmark Dataset: UCI Cleveland Cohort", level=2)
    add_paragraph_styled(
        doc,
        "The historical benchmark utilized is the canonical Cleveland Clinic Foundation heart disease database (N=303, 13 predictors, "
        "Janosi et al., 1988). The UCI dataset contains nuclear thallium scintigraphy ('thal'), which is absent in the Indian cohort, "
        "and codes chest pain from 1 to 4 rather than 0 to 3. These structural incompatibilities reinforce the requirement for separate model pipelines."
    )

    # ==========================================
    # SECTION 3: EXPLORATORY DATA ANALYSIS
    # ==========================================
    add_header_styled(doc, "3. Exploratory Data Analysis & Visualization Suite", level=1)
    add_paragraph_styled(
        doc,
        "A complete suite of publication-grade figures was generated for the Indian hospital cohort to analyze marginal distributions, "
        "inter-feature correlations, and bivariate clinical risk separations."
    )

    add_image_with_caption(doc, "target_distribution.png", "Target Class Distribution in Indian Hospital Cohort (58% Positive vs. 42% Negative)", cohort="india")
    add_image_with_caption(doc, "numeric_distributions.png", "Continuous Feature Distributions with Kernel Density Estimates (Age, BP, Cholesterol, Max HR, Oldpeak)", cohort="india")
    add_image_with_caption(doc, "numeric_boxplots.png", "Outlier Diagnostics across Continuous Clinical Attributes", cohort="india")
    add_image_with_caption(doc, "correlation_heatmap.png", "Pearson Correlation Heatmap Demonstrating Feature Orthogonality and Absence of Multicollinearity", cohort="india")
    add_image_with_caption(doc, "categorical_distributions.png", "Categorical Proportions across Clinical Diagnostic Subcategories", cohort="india")
    add_image_with_caption(doc, "feature_target_relationships.png", "Bivariate Clinical Risk Relationships: Feature Distributions Stratified by Cardiovascular Disease Status", cohort="india")

    # ==========================================
    # SECTION 4: LEAKAGE-SAFE PREPROCESSING PIPELINE
    # ==========================================
    add_header_styled(doc, "4. Leakage-Safe Preprocessing & ColumnTransformer Architecture", level=1)
    add_paragraph_styled(
        doc,
        "A critical vulnerability in clinical machine learning is data leakage—where information from the evaluation split inadvertently "
        "contaminates preprocessing parameters. CardioSense strictly enforces a Split-First protocol using Scikit-Learn ColumnTransformer and Pipeline."
    )

    add_formula_box(
        doc,
        "Z-Score Standardization (StandardScaler)",
        "z_ij = (x_ij - μ_j(train)) / σ_j(train)",
        "Where μ_j and σ_j are computed strictly on the training partition and applied identically to test partitions without recalculation."
    )

    preproc_steps = [
        ("Stratified Split Isolation: ", "80% training (N=800) and 20% held-out test (N=200) partitioned with fixed random seed (R=42) maintaining exact class prevalence."),
        ("Numeric Transformer: ", "SimpleImputer(strategy='median') followed by StandardScaler(). The median is robust to acute outliers in blood pressure and cholesterol."),
        ("Categorical Transformer: ", "SimpleImputer(strategy='most_frequent') followed by OneHotEncoder(handle_unknown='ignore', sparse_output=False) to prevent crashes on unseen clinical categories."),
        ("Pipeline Encapsulation: ", "The entire preprocessing transformer and estimator are bundled into a single serializable Pipeline object. The raw test data passes directly into pipeline.predict().")
    ]
    for bold_pre, text in preproc_steps:
        add_paragraph_styled(doc, text=text, bold_prefix="• " + bold_pre, space_after=3)

    # ==========================================
    # SECTION 5: MODEL ARCHITECTURES & MATHEMATICAL FORMULATIONS
    # ==========================================
    add_header_styled(doc, "5. Evaluated Model Architectures & Mathematical Foundations", level=1)
    
    add_paragraph_styled(
        doc,
        "Three distinct algorithm families were systematically evaluated to assess the bias-variance trade-off across linear and non-linear inductive biases."
    )

    add_header_styled(doc, "5.1 Logistic Regression with L2 Regularization", level=2)
    add_formula_box(
        doc,
        "L2-Regularized Logistic Objective (Ridge Logit)",
        "min_w  - ∑ [ y_i log(σ(w^T x_i)) + (1 - y_i) log(1 - σ(w^T x_i)) ] + (1 / 2C) ||w||_2^2",
        "Where σ(z) = 1 / (1 + e^(-z)) and C is the inverse regularization parameter tuned via 5-fold CV."
    )

    add_header_styled(doc, "5.2 Random Forest Classifier", level=2)
    add_formula_box(
        doc,
        "Random Forest Bagging & Variance Reduction",
        "Var(f_RF(x)) = ρ σ^2 + ((1 - ρ) / B) σ^2",
        "Where B is the number of bootstrap trees and ρ is the pairwise tree correlation reduced via random feature subspace sampling."
    )

    add_header_styled(doc, "5.3 XGBoost (Extreme Gradient Boosting)", level=2)
    add_formula_box(
        doc,
        "XGBoost Second-Order Taylor Objective",
        "Obj^(t) ≈ ∑ [ g_i f_t(x_i) + (1/2) h_i f_t^2(x_i) ] + γ T + (λ/2) ∑ w_j^2",
        "Where g_i = ∂l/∂ŷ and h_i = ∂^2l/∂ŷ^2 are first and second order loss gradients, γ penalizes leaf count T, and λ penalizes leaf weights w_j."
    )

    # ==========================================
    # SECTION 6: REAL MEASURED EXPERIMENTAL RESULTS
    # ==========================================
    add_header_styled(doc, "6. Real Measured Experimental Results & Dual-Cohort Benchmark", level=1)
    
    add_paragraph_styled(
        doc,
        "Experiments were executed using Stratified 5-Fold Cross-Validation on training data and evaluated on untouched held-out test splits. "
        "No metrics are fabricated; all values represent actual measured empirical outputs."
    )

    add_header_styled(doc, "6.1 Primary Indian Hospital Cohort Benchmark (N_train=800, N_test=200)", level=2)
    india_res_headers = ["Model Family", "5-Fold CV ROC-AUC", "Test Accuracy", "Test Recall (Sens.)", "Test Precision", "Test F1", "Test ROC-AUC", "False Neg.", "False Pos."]
    india_res_data = [
        ["Logistic Regression", "0.9909 ± 0.0084", "97.00%", "97.41%", "97.41%", "0.9741", "0.9978", "3", "3"],
        ["Random Forest", "0.9957 ± 0.0059", "99.00%", "99.14%", "99.14%", "0.9914", "0.9995", "1", "1"],
        ["XGBoost (Champion)", "0.9960 ± 0.0040", "99.50%", "100.0%", "99.15%", "0.9957", "0.9998", "0", "1"]
    ]
    create_styled_table(doc, india_res_headers, india_res_data, [1.4, 1.1, 0.7, 0.8, 0.7, 0.6, 0.7, 0.5, 0.5])

    add_callout(
        doc,
        "CRITICAL CLINICAL FINDING: ZERO FALSE NEGATIVES",
        "On the held-out test split of 200 Indian patients (116 confirmed CVD cases), the champion XGBoost architecture achieved "
        "a perfect 100.0% sensitivity (Recall = 1.0000, False Negatives = 0). In cardiovascular screening, false negatives represent "
        "undiagnosed patients who may suffer fatal infarctions. Eliminating false negatives while maintaining 99.15% precision "
        "demonstrates exceptional discriminative utility on this hospital cohort.",
        icon="🎯"
    )

    add_header_styled(doc, "6.2 UCI Cleveland Benchmark Cohort (N_train=242, N_test=61)", level=2)
    uci_res_headers = ["Model Family", "5-Fold CV ROC-AUC", "Test Accuracy", "Test Recall", "Test Precision", "Test F1", "Test ROC-AUC", "False Neg.", "False Pos."]
    uci_res_data = [
        ["Logistic Regression", "0.9025 ± 0.0144", "88.52%", "92.86%", "83.87%", "0.8814", "0.9665", "2", "5"],
        ["Random Forest (Champion)", "0.9038 ± 0.0238", "88.52%", "92.86%", "83.87%", "0.8814", "0.9502", "2", "5"],
        ["XGBoost", "0.8901 ± 0.0221", "90.16%", "92.86%", "86.67%", "0.8966", "0.9481", "2", "4"]
    ]
    create_styled_table(doc, uci_res_headers, uci_res_data, [1.4, 1.1, 0.7, 0.8, 0.7, 0.6, 0.7, 0.5, 0.5])

    add_image_with_caption(doc, "confusion_matrices_comparison.png", "Confusion Matrix Comparison across Architectures on Indian Cohort (XGBoost achieves 0 FN)", cohort="india")
    add_image_with_caption(doc, "roc_curves_comparison.png", "Receiver Operating Characteristic (ROC) Curves for Indian Cohort", cohort="india")
    add_image_with_caption(doc, "precision_recall_curves_comparison.png", "Precision-Recall Curves Demonstrating High AUPRC", cohort="india")
    add_image_with_caption(doc, "models_metrics_comparison.png", "Multi-Metric Bar Chart Comparison across Algorithms", cohort="india")

    # ==========================================
    # SECTION 7: SHAP EXPLAINABILITY
    # ==========================================
    add_header_styled(doc, "7. Game-Theoretic SHAP Feature Attribution Analysis", level=1)
    add_paragraph_styled(
        doc,
        "CardioSense incorporates Shapley Additive exPlanations (SHAP) grounded in cooperative game theory. "
        "Unlike heuristic tree feature importance (Gini impurity decrease), which is biased toward high-cardinality features, "
        "Shapley values satisfy efficiency, symmetry, dummy player, and additivity axioms."
    )

    add_formula_box(
        doc,
        "Shapley Value Formulation",
        "ϕ_i(v) = ∑_(S ⊆ N \\ {i}) [ |S|! (|N| - |S| - 1)! / |N|! ] * [ v(S ∪ {i}) - v(S) ]",
        "Computing the marginal contribution of clinical feature i averaged over all possible feature coalitions S."
    )

    add_image_with_caption(doc, "shap_summary_plot.png", "SHAP Beeswarm Summary Plot: Feature Distributions and Directional Impact on Disease Risk", cohort="india")
    add_image_with_caption(doc, "shap_feature_importance_bar.png", "Mean Absolute SHAP Feature Importance Ranking for Indian Hospital Cohort", cohort="india")
    add_image_with_caption(doc, "shap_local_correct.png", "Local Patient Waterfall Explanation: Correctly Classified High-Risk Cardiac Patient", cohort="india")

    add_header_styled(doc, "7.1 Top Clinical Attribution Rankings (Indian Cohort)", level=2)
    shap_headers = ["Rank", "Predictor Feature", "Mean |SHAP Value|", "Clinical Interpretation & Directional Effect"]
    shap_data = [
        ["1", "chestpain (Chest Pain Type)", "1.341", "Highest predictive contribution; asymptomatic presentation or severe typical angina strongly shifts probability toward CAD presence."],
        ["2", "oldpeak (ST Depression)", "0.985", "Exercise-induced ST depression reflects severe subendocardial ischemia; higher oldpeak dramatically elevates risk."],
        ["3", "noofmajorvessels (Vessels Colored)", "0.762", "Fluoroscopy-detected vessel narrowing directly corresponds to multi-vessel CAD."],
        ["4", "maxheartrate (Peak Exercise HR)", "0.612", "Chronotropic incompetence (inability to reach expected peak HR) strongly correlates with cardiac dysfunction."],
        ["5", "exerciseangia (Exercise Angina)", "0.448", "Angina provoked by exertion confirms hemodynamic inadequacy under cardiac workload."]
    ]
    create_styled_table(doc, shap_headers, shap_data, [0.6, 1.8, 1.1, 3.0])

    # ==========================================
    # SECTION 8: CLINICAL ERROR DIAGNOSTICS
    # ==========================================
    add_header_styled(doc, "8. Clinical Error Diagnostics & Misclassification Analysis", level=1)
    add_paragraph_styled(
        doc,
        "Rigorous data science requires auditing model errors. On the Indian hospital test set (N=200), the champion XGBoost model "
        "produced exactly 1 misclassification out of 200 cases (0 False Negatives, 1 False Positive)."
    )

    error_headers = ["Metric / Case", "Observed Count", "Clinical Risk Level", "Diagnostic Remediation Strategy"]
    error_data = [
        ["False Negatives (FN)", "0 / 116 CVD Cases", "Zero Clinical Danger", "No diseased patients missed; exceptional safety profile for primary triage."],
        ["False Positives (FP)", "1 / 84 Healthy Cases", "Low Impact (False Alarm)", "One healthy patient referred for secondary non-invasive testing (Echo/Stress test)."],
        ["Sensitivity (Recall)", "100.0%", "Gold Standard", "Meets strict requirements for high-stakes medical screening tools."],
        ["Specificity", "98.81%", "High Precision", "Minimal patient anxiety and near-zero unnecessary catheterization referrals."]
    ]
    create_styled_table(doc, error_headers, error_data, [1.5, 1.2, 1.3, 2.5])

    # ==========================================
    # SECTION 9: PRODUCTION MLOPS & SYSTEM ARCHITECTURE
    # ==========================================
    add_header_styled(doc, "9. Production MLOps, System Architecture & User Interfaces", level=1)
    add_paragraph_styled(
        doc,
        "CardioSense features a modular production architecture supporting dual-cohort inference without code duplication:"
    )

    arch_points = [
        ("FastAPI Backend Service: ", "Exposes dedicated endpoints /predict/india (12 features) and /predict/uci (13 features), alongside a dynamic /predict endpoint that automatically validates payloads against the appropriate Pydantic schema."),
        ("Strict Schema Validation: ", "Pydantic models reject physiological impossibilities (e.g., negative age, systolic BP outside 80-250 mm Hg) with HTTP 422 Unprocessable Entity."),
        ("Multi-Model Lifespan Loading: ", "On application startup, pre-trained pipelines for both cohorts are deserialized into memory without re-training."),
        ("Full-Stack Next.js 16 Web Portal: ", "frontend/ provides an interactive React 19 web portal styled with the soft pastel design system (ivory #FAF8F4, sage #E8ECE9, mint #D8E2DC, lavender #E8E5EE, peach #FDE2D2, forest green #355B4B accents). Features a dual-cohort switcher, live patient risk simulation, benchmark tables, and interactive SHAP lightbox.")
    ]
    for bold_pre, text in arch_points:
        add_paragraph_styled(doc, text=text, bold_prefix="• " + bold_pre, space_after=3)

    # ==========================================
    # SECTION 10: LIMITATIONS & ETHICAL GOVERNANCE
    # ==========================================
    add_header_styled(doc, "10. Scientific Limitations & Ethical Governance", level=1)
    limitations = [
        ("Single-Center Hospital Inpatient Sample: ", "The Indian cohort represents patients referred to a tertiary care hospital for cardiac evaluation, yielding a baseline disease prevalence (58.0%) higher than routine community screening."),
        ("Geographic Generalizability across India: ", "India is genetically and socio-economically diverse. As demonstrated by ICMR-INDIAB, cardiovascular prevalence varies markedly between northern and southern states. Multi-center validation is required before generalized deployment."),
        ("Surrogate vs. Definitive Angiography: ", "Machine learning outputs provide statistical risk estimates, not anatomical visualization. Clinicians must not use predictions to cancel angiography when clinical suspicion is high."),
        ("Absence of Modern Biomarkers: ", "Neither dataset includes high-sensitivity cardiac troponin I/T (hs-cTn), coronary computed tomography angiography (CCTA), or coronary artery calcium (CAC) scoring.")
    ]
    for bold_pre, text in limitations:
        add_paragraph_styled(doc, text=text, bold_prefix="• " + bold_pre, space_after=4)

    # ==========================================
    # SECTION 11: REFERENCES
    # ==========================================
    add_header_styled(doc, "11. Primary Academic References", level=1)
    refs = [
        "1. Doppala, K., & Bhattacharyya, D. (2020). Cardiovascular Disease Dataset. Mendeley Data, V1, DOI: 10.17632/dzz48mvjht.1.",
        "2. Anjana, R. M., et al. (ICMR-INDIAB Collaborative Study Group). (2023). Metabolic non-communicable disease health report in India. The Lancet Diabetes & Endocrinology, 11(7), 474–489.",
        "3. Janosi, A., Steinbrunn, W., Pfisterer, M., & Detrano, R. (1988). Heart Disease Database. Cleveland Clinic Foundation. UCI Machine Learning Repository.",
        "4. Lundberg, S. M., & Lee, S. I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems (NeurIPS 2017), 30, 4765–4774.",
        "5. Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 785–794.",
        "6. Breiman, L. (2001). Random forests. Machine Learning, 45(1), 5–32.",
        "7. World Health Organization (WHO). (2024). Cardiovascular Diseases Fact Sheet. Geneva: WHO."
    ]
    for r in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(3)
        p_ref.paragraph_format.line_spacing = 1.15
        r_run = p_ref.add_run(r)
        r_run.font.name = "Calibri"
        r_run.font.size = Pt(9.5)
        r_run.font.color.rgb = TEXT_DARK

    # Save document
    doc.save(str(OUTPUT_DOCX))
    print(f"Successfully generated comprehensive Data Science document at: {OUTPUT_DOCX}")

if __name__ == "__main__":
    build_document()
